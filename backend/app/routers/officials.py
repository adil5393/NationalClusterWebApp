"""CBSE officials monitoring the championship — a flat roster (see
models.Official), independent of any participating Team. Bundles everything
official-specific in one router: CRUD, admin photo upload, bulk Excel
import + a downloadable example template for it, and ID card
generation/download (single, and "download all" as an A4 sheet, a 12x18in
print-shop sheet, or a ZIP of individual card PDFs) — mirroring the
participant ID card exports in exports.py, just against id_card.py's
official-specific render functions and without any team/age-group
grouping, since officials are a single flat list.
"""
import io
import uuid
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from .. import id_card, models, schemas
from ..database import get_db
from ..image_utils import optimize_image

router = APIRouter(prefix="/api/officials", tags=["officials"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

ASSETS_OFFICIALS_DIR = Path(__file__).resolve().parent.parent.parent / "assets" / "officials"
VALID_PHOTO_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


# --- CRUD ----------------------------------------------------------------
@router.get("", response_model=list[schemas.OfficialRead])
def list_officials(db: Session = Depends(get_db)):
    return db.query(models.Official).order_by(models.Official.full_name).all()


@router.post("", response_model=schemas.OfficialRead, status_code=201)
def create_official(payload: schemas.OfficialCreate, db: Session = Depends(get_db)):
    v = models.Official(**payload.model_dump())
    db.add(v)
    db.commit()
    db.refresh(v)
    return v


@router.put("/{official_id}", response_model=schemas.OfficialRead)
def update_official(official_id: int, payload: schemas.OfficialUpdate, db: Session = Depends(get_db)):
    v = db.get(models.Official, official_id)
    if not v:
        raise HTTPException(404, "Official not found")
    for k, val in payload.model_dump(exclude_unset=True).items():
        setattr(v, k, val)
    db.commit()
    db.refresh(v)
    return v


@router.delete("/{official_id}", status_code=204)
def delete_official(official_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Official, official_id)
    if not v:
        raise HTTPException(404, "Official not found")
    if v.photo_filename:
        path = ASSETS_OFFICIALS_DIR / v.photo_filename
        if path.exists() and path.is_file():
            path.unlink()
    db.delete(v)
    db.commit()


# --- Photo (admin-side upload, like gallery.py) ---------------------------
@router.post("/{official_id}/photo", response_model=schemas.OfficialRead)
def upload_official_photo(official_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    v = db.get(models.Official, official_id)
    if not v:
        raise HTTPException(404, "Official not found")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in VALID_PHOTO_EXTENSIONS:
        raise HTTPException(400, "Unsupported file type (use JPG, PNG, or WEBP)")

    ASSETS_OFFICIALS_DIR.mkdir(parents=True, exist_ok=True)
    content = optimize_image(file.file.read(), ext)
    name = f"{uuid.uuid4().hex}{ext}"
    (ASSETS_OFFICIALS_DIR / name).write_bytes(content)

    old_filename = v.photo_filename
    v.photo_filename = name
    db.commit()
    db.refresh(v)
    if old_filename:
        old_path = ASSETS_OFFICIALS_DIR / old_filename
        if old_path.exists() and old_path.is_file():
            old_path.unlink()
    return v


@router.delete("/{official_id}/photo", response_model=schemas.OfficialRead)
def delete_official_photo(official_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Official, official_id)
    if not v:
        raise HTTPException(404, "Official not found")
    if v.photo_filename:
        path = ASSETS_OFFICIALS_DIR / v.photo_filename
        if path.exists() and path.is_file():
            path.unlink()
        v.photo_filename = None
        db.commit()
        db.refresh(v)
    return v


# --- Excel import + example template --------------------------------------
IMPORT_COLUMNS = ["full_name", "designation", "organization", "official_id_no", "gender", "phone", "email", "notes"]


def _val(row: dict, key: str) -> "str | None":
    v = row.get(key)
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # numeric cells (phones, ids) must not become "9876543210.0"
    s = str(v).strip()
    return s or None


@router.get("/template.xlsx")
def download_official_template():
    """A blank example spreadsheet with the exact headers /import expects,
    plus one filled-in sample row — so whoever fills it out sees the
    expected shape rather than guessing at column names."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Officials"
    headers = ["full_name", "designation", "organization", "official_id_no", "gender", "phone", "email", "notes"]
    sample = ["Rakesh Kumar", "Observer", "CBSE Regional Office, Lucknow", "CBSE/OBS/014", "Male", "9876543210", "rakesh@example.com", "Zonal monitoring"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    for col, val in enumerate(sample, start=1):
        ws.cell(row=2, column=col, value=val)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="officials_template.xlsx"'},
    )


@router.post("/import")
async def import_officials(file: UploadFile = File(...), db: Session = Depends(get_db)):
    content = await file.read()
    name = (file.filename or "").lower()
    try:
        if name.endswith(".xlsx") or name.endswith(".xls"):
            df = pd.read_excel(io.BytesIO(content))
        else:
            df = pd.read_csv(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Could not parse file: {e}")
    df.columns = [str(c).strip().lower() for c in df.columns]
    rows = df.to_dict("records")

    existing_names = {v.full_name.strip().lower() for v in db.query(models.Official).all()}
    created, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):
        full_name = _val(row, "full_name") or _val(row, "name")
        if not full_name:
            errors.append(f"Row {i}: missing 'full_name'")
            continue
        if full_name.strip().lower() in existing_names:
            skipped += 1
            continue
        db.add(models.Official(
            full_name=full_name,
            designation=_val(row, "designation") or _val(row, "title"),
            organization=_val(row, "organization"),
            official_id_no=_val(row, "official_id_no") or _val(row, "official_id"),
            gender=_val(row, "gender"),
            phone=_val(row, "phone"),
            email=_val(row, "email"),
            notes=_val(row, "notes"),
        ))
        existing_names.add(full_name.strip().lower())
        created += 1
    db.commit()
    return {"entity": "officials", "created": created, "skipped": skipped, "errors": errors}


# --- ID cards ---------------------------------------------------------
def _photo_path(official: models.Official) -> "Path | None":
    if not official.photo_filename:
        return None
    return ASSETS_OFFICIALS_DIR / official.photo_filename


def _pdf_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter([content]),
        media_type=PDF_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _zip_response(files: list[tuple[str, bytes]], filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in files:
            zf.writestr(name, content)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _idcard_filename(official: models.Official) -> str:
    slug = "".join(c if c.isalnum() or c in " -_" else "_" for c in official.full_name).strip() or "official"
    return f"{official.id}_{slug}.pdf"


@router.get("/{official_id}/idcard.pdf")
def export_official_idcard(official_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Official, official_id)
    if not v:
        raise HTTPException(404, "Official not found")
    card = id_card.render_official_id_card_page(v, _photo_path(v))
    pdf = id_card.build_pdf([card])
    return _pdf_response(pdf, f"official-idcard-{v.id}.pdf")


@router.get("/idcards/all.pdf")
def export_all_official_idcards(db: Session = Depends(get_db)):
    officials = db.query(models.Official).order_by(models.Official.full_name).all()
    if not officials:
        raise HTTPException(404, "No officials to generate cards for")
    cards = [id_card.render_official_id_card(v, _photo_path(v)) for v in officials]
    pdf = id_card.build_pdf_sheets([cards], layout=id_card.staff_sheet_layout(id_card.A4_SHEET), dpi=id_card.PRINT_DPI)
    return _pdf_response(pdf, "official-idcards-all.pdf")


@router.get("/idcards/all/sheet-12x18.pdf")
def export_all_official_idcards_12x18(db: Session = Depends(get_db)):
    officials = db.query(models.Official).order_by(models.Official.full_name).all()
    if not officials:
        raise HTTPException(404, "No officials to generate cards for")
    cards = [id_card.render_official_id_card(v, _photo_path(v)) for v in officials]
    pdf = id_card.build_pdf_sheets([cards], layout=id_card.staff_sheet_layout(id_card.SHEET_12X18), dpi=id_card.PRINT_DPI)
    return _pdf_response(pdf, "official-idcards-all-12x18.pdf")


@router.get("/idcards/all/individual.zip")
def export_all_official_idcards_individual(db: Session = Depends(get_db)):
    officials = db.query(models.Official).order_by(models.Official.full_name).all()
    if not officials:
        raise HTTPException(404, "No officials to generate cards for")
    files: list[tuple[str, bytes]] = []
    seen_names: dict[str, int] = {}
    for v in officials:
        card = id_card.render_official_id_card_page(v, _photo_path(v))
        pdf = id_card.build_pdf([card])
        name = _idcard_filename(v)
        if name in seen_names:
            seen_names[name] += 1
            stem, _, ext = name.rpartition(".")
            name = f"{stem}_{seen_names[name]}.{ext}"
        else:
            seen_names[name] = 0
        files.append((name, pdf))
    return _zip_response(files, "official-idcards-all-individual.zip")
