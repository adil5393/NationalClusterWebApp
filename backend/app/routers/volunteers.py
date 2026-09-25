"""The host school's own event volunteers — a flat roster (see
models.Volunteer), independent of any participating Team. Bundles everything
volunteer-specific in one router: CRUD, admin photo upload, bulk Excel
import + a downloadable example template for it, and ID card
generation/download (single, and "download all" as an A4 sheet, a 12x18in
print-shop sheet, or a ZIP of individual card PDFs) — mirroring the
participant ID card exports in exports.py, just against id_card.py's
volunteer-specific render functions and without any team/age-group
grouping, since volunteers are a single flat list.
"""
import io
import uuid
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from .. import id_card, models, schemas
from ..auth_utils import hash_password, provision_login_credentials
from ..database import get_db
from ..image_utils import optimize_image

router = APIRouter(prefix="/api/volunteers", tags=["volunteers"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

ASSETS_VOLUNTEERS_DIR = Path(__file__).resolve().parent.parent.parent / "assets" / "volunteers"
VALID_PHOTO_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


# --- CRUD ----------------------------------------------------------------
@router.get("", response_model=list[schemas.VolunteerRead])
def list_volunteers(db: Session = Depends(get_db)):
    return db.query(models.Volunteer).order_by(models.Volunteer.full_name).all()


@router.post("", response_model=schemas.VolunteerRead, status_code=201)
def create_volunteer(payload: schemas.VolunteerCreate, db: Session = Depends(get_db)):
    v = models.Volunteer(**payload.model_dump())
    db.add(v)
    db.commit()
    db.refresh(v)
    return v


@router.put("/{volunteer_id}", response_model=schemas.VolunteerRead)
def update_volunteer(volunteer_id: int, payload: schemas.VolunteerUpdate, db: Session = Depends(get_db)):
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    for k, val in payload.model_dump(exclude_unset=True).items():
        setattr(v, k, val)
    db.commit()
    db.refresh(v)
    return v


@router.delete("/{volunteer_id}", status_code=204)
def delete_volunteer(volunteer_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    if v.photo_filename:
        path = ASSETS_VOLUNTEERS_DIR / v.photo_filename
        if path.exists() and path.is_file():
            path.unlink()
    db.delete(v)
    db.commit()


@router.post("/{volunteer_id}/credential", response_model=schemas.VolunteerCredentialResult, status_code=201)
def create_volunteer_credential(volunteer_id: int, db: Session = Depends(get_db)):
    """Provisions a self-service Organizer Portal login for one volunteer —
    the volunteer counterpart to routers/staff.py's create_staff_credential.
    Grants schemas.VOLUNTEER_BASE_PERMISSIONS (no organizer module access);
    the account can only reach its own profile/ID card via /me/volunteer."""
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    if v.organizer_users:
        raise HTTPException(409, f"{v.full_name} already has a login: {v.organizer_users[0].username}")

    username, password = provision_login_credentials(db, v.full_name, fallback_label="VOLUNTEER")
    login = models.OrganizerUser(
        username=username,
        full_name=v.full_name,
        password_hash=hash_password(password),
        is_active=True,
        is_admin=False,
        permissions=schemas.VOLUNTEER_BASE_PERMISSIONS,
        volunteers=[v],
    )
    db.add(login)
    db.commit()
    return {"login_username": username, "login_password": password}


# --- Photo (admin-side upload, like gallery.py) ---------------------------
@router.post("/{volunteer_id}/photo", response_model=schemas.VolunteerRead)
def upload_volunteer_photo(volunteer_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in VALID_PHOTO_EXTENSIONS:
        raise HTTPException(400, "Unsupported file type (use JPG, PNG, or WEBP)")

    ASSETS_VOLUNTEERS_DIR.mkdir(parents=True, exist_ok=True)
    content = optimize_image(file.file.read(), ext)
    name = f"{uuid.uuid4().hex}{ext}"
    (ASSETS_VOLUNTEERS_DIR / name).write_bytes(content)

    old_filename = v.photo_filename
    v.photo_filename = name
    db.commit()
    db.refresh(v)
    if old_filename:
        old_path = ASSETS_VOLUNTEERS_DIR / old_filename
        if old_path.exists() and old_path.is_file():
            old_path.unlink()
    return v


@router.delete("/{volunteer_id}/photo", response_model=schemas.VolunteerRead)
def delete_volunteer_photo(volunteer_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    if v.photo_filename:
        path = ASSETS_VOLUNTEERS_DIR / v.photo_filename
        if path.exists() and path.is_file():
            path.unlink()
        v.photo_filename = None
        db.commit()
        db.refresh(v)
    return v


# --- Excel import + example template --------------------------------------
IMPORT_COLUMNS = ["full_name", "student_class", "gender", "phone", "email", "notes"]


def _val(row: dict, key: str) -> "str | None":
    v = row.get(key)
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # numeric cells (phones, ids) must not become "9876543210.0"
    s = str(v).strip()
    return s or None


@router.get("/template.xlsx")
def download_volunteer_template():
    """A blank example spreadsheet with the exact headers /import expects,
    plus one filled-in sample row — so whoever fills it out sees the
    expected shape rather than guessing at column names."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Volunteers"
    headers = ["full_name", "student_class", "gender", "phone", "email", "notes"]
    sample = ["Ananya Sharma", "XII", "Female", "9876543210", "ananya@example.com", "Registration desk"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    for col, val in enumerate(sample, start=1):
        ws.cell(row=2, column=col, value=val)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="volunteers_template.xlsx"'},
    )


@router.post("/import")
async def import_volunteers(file: UploadFile = File(...), db: Session = Depends(get_db)):
    content = await file.read()
    name = (file.filename or "").lower()
    try:
        if name.endswith(".xlsx") or name.endswith(".xls"):
            df = pd.read_excel(io.BytesIO(content))
        else:
            df = pd.read_csv(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Could not parse file: {e}")
    df.columns = [str(c).strip().lower() for c in df.columns]
    rows = df.to_dict("records")

    existing_names = {v.full_name.strip().lower() for v in db.query(models.Volunteer).all()}
    created, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):
        full_name = _val(row, "full_name") or _val(row, "name")
        if not full_name:
            errors.append(f"Row {i}: missing 'full_name'")
            continue
        if full_name.strip().lower() in existing_names:
            skipped += 1
            continue
        db.add(models.Volunteer(
            full_name=full_name,
            student_class=_val(row, "student_class") or _val(row, "class"),
            gender=_val(row, "gender"),
            phone=_val(row, "phone"),
            email=_val(row, "email"),
            notes=_val(row, "notes"),
        ))
        existing_names.add(full_name.strip().lower())
        created += 1
    db.commit()
    return {"entity": "volunteers", "created": created, "skipped": skipped, "errors": errors}


# --- ID cards ---------------------------------------------------------
def _photo_path(volunteer: models.Volunteer) -> "Path | None":
    if not volunteer.photo_filename:
        return None
    return ASSETS_VOLUNTEERS_DIR / volunteer.photo_filename


def _pdf_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter([content]),
        media_type=PDF_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _zip_response(files: list[tuple[str, bytes]], filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in files:
            zf.writestr(name, content)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _idcard_filename(volunteer: models.Volunteer) -> str:
    slug = "".join(c if c.isalnum() or c in " -_" else "_" for c in volunteer.full_name).strip() or "volunteer"
    return f"{volunteer.id}_{slug}.pdf"


@router.get("/{volunteer_id}/idcard.pdf")
def export_volunteer_idcard(volunteer_id: int, db: Session = Depends(get_db)):
    v = db.get(models.Volunteer, volunteer_id)
    if not v:
        raise HTTPException(404, "Volunteer not found")
    card = id_card.render_volunteer_id_card_page(v, _photo_path(v))
    pdf = id_card.build_pdf([card])
    return _pdf_response(pdf, f"volunteer-idcard-{v.id}.pdf")


@router.get("/idcards/all.pdf")
def export_all_volunteer_idcards(db: Session = Depends(get_db)):
    volunteers = db.query(models.Volunteer).order_by(models.Volunteer.full_name).all()
    if not volunteers:
        raise HTTPException(404, "No volunteers to generate cards for")
    cards = [id_card.render_volunteer_id_card(v, _photo_path(v)) for v in volunteers]
    pdf = id_card.build_pdf_sheets([cards], layout=id_card.A4_SHEET, dpi=id_card.PRINT_DPI)
    return _pdf_response(pdf, "volunteer-idcards-all.pdf")


@router.get("/idcards/all/sheet-12x18.pdf")
def export_all_volunteer_idcards_12x18(db: Session = Depends(get_db)):
    volunteers = db.query(models.Volunteer).order_by(models.Volunteer.full_name).all()
    if not volunteers:
        raise HTTPException(404, "No volunteers to generate cards for")
    cards = [id_card.render_volunteer_id_card(v, _photo_path(v)) for v in volunteers]
    pdf = id_card.build_pdf_sheets([cards], layout=id_card.SHEET_12X18, dpi=id_card.PRINT_DPI)
    return _pdf_response(pdf, "volunteer-idcards-all-12x18.pdf")


@router.get("/idcards/all/individual.zip")
def export_all_volunteer_idcards_individual(db: Session = Depends(get_db)):
    volunteers = db.query(models.Volunteer).order_by(models.Volunteer.full_name).all()
    if not volunteers:
        raise HTTPException(404, "No volunteers to generate cards for")
    files: list[tuple[str, bytes]] = []
    seen_names: dict[str, int] = {}
    for v in volunteers:
        card = id_card.render_volunteer_id_card_page(v, _photo_path(v))
        pdf = id_card.build_pdf([card])
        name = _idcard_filename(v)
        if name in seen_names:
            seen_names[name] += 1
            stem, _, ext = name.rpartition(".")
            name = f"{stem}_{seen_names[name]}.{ext}"
        else:
            seen_names[name] = 0
        files.append((name, pdf))
    return _zip_response(files, "volunteer-idcards-all-individual.zip")
